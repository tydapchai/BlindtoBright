import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import re

doc = docx.Document()

# Margin chuẩn
for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

# Màu sắc
COLOR_TITLE = RGBColor(15, 23, 42)
COLOR_H1 = RGBColor(26, 54, 93)
COLOR_H2 = RGBColor(30, 64, 175)
COLOR_TEXT = RGBColor(30, 41, 59)

def add_styled_heading(text, level):
    h = doc.add_heading(level=level)
    run = h.add_run(text)
    run.font.name = "Arial"
    if level == 1:
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = COLOR_H1
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(6)
    elif level == 2:
        run.font.size = Pt(13.5)
        run.font.bold = True
        run.font.color.rgb = COLOR_H2
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
    elif level == 3:
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = COLOR_TEXT
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(2)
    return h

def add_p(text, bold_prefix=None, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.25
    if bold_prefix:
        r_bold = p.add_run(bold_prefix)
        r_bold.font.name = "Arial"
        r_bold.font.size = Pt(11)
        r_bold.font.bold = True
        r_bold.font.color.rgb = COLOR_TEXT
    
    # Simple parse for **bold** text in markdown
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            r = p.add_run(part[2:-2])
            r.font.bold = True
        else:
            r = p.add_run(part)
        r.font.name = "Arial"
        r.font.size = Pt(11)
        r.font.color.rgb = COLOR_TEXT
    return p

def add_bullet(text, space_after=3):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.2
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            r = p.add_run(part[2:-2])
            r.font.bold = True
        else:
            r = p.add_run(part)
        r.font.name = "Arial"
        r.font.size = Pt(11)
        r.font.color.rgb = COLOR_TEXT
    return p

# Đọc file markdown
with open(r"d:\BlindtoBright\docs\TAI_LIEU_DU_AN_BLINDTOBRIGHT.md", "r", encoding="utf-8") as f:
    lines = f.readlines()

in_table = False
table_rows = []

for line in lines:
    line_s = line.strip()
    
    if line_s.startswith("|") and line_s.endswith("|"):
        # Header separator row like |:---|:---|
        if ":---" in line_s or "---:" in line_s or "---" in line_s and not any(c.isalnum() for c in line_s):
            continue
        row_items = [c.strip() for c in line_s[1:-1].split("|")]
        table_rows.append(row_items)
        in_table = True
        continue
    else:
        if in_table and table_rows:
            # Tạo bảng Word
            cols = max(len(r) for r in table_rows)
            t = doc.add_table(rows=len(table_rows), cols=cols)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for r_idx, r_data in enumerate(table_rows):
                for c_idx, cell_value in enumerate(r_data):
                    cell = t.cell(r_idx, c_idx)
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.space_before = Pt(2)
                    parts = re.split(r'(\*\*.*?\*\*)', cell_value)
                    for part in parts:
                        if part.startswith('**') and part.endswith('**'):
                            r = p.add_run(part[2:-2])
                            r.font.bold = True
                        else:
                            r = p.add_run(part)
                        r.font.name = "Arial"
                        r.font.size = Pt(10)
                        if r_idx == 0:
                            r.font.bold = True
                    # Shading for header
                    if r_idx == 0:
                        tcPr = cell._tc.get_or_add_tcPr()
                        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>')
                        tcPr.append(shd)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            table_rows = []
            in_table = False

    if not line_s:
        continue
    
    if line_s.startswith("<!--") and line_s.endswith("-->"):
        continue

    if line_s.startswith("# "):
        add_styled_heading(line_s[2:], level=1)
    elif line_s.startswith("## "):
        add_styled_heading(line_s[3:], level=1)
    elif line_s.startswith("### "):
        add_styled_heading(line_s[4:], level=2)
    elif line_s.startswith("#### "):
        add_styled_heading(line_s[5:], level=3)
    elif line_s.startswith("* ") or line_s.startswith("- "):
        add_bullet(line_s[2:])
    elif re.match(r'^\d+\.\s', line_s):
        # Numbered list
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        num_str = line_s.split('. ', 1)[0] + '. '
        content = line_s.split('. ', 1)[1]
        r_num = p.add_run(num_str)
        r_num.font.bold = True
        r_num.font.name = "Arial"
        r_num.font.size = Pt(11)
        parts = re.split(r'(\*\*.*?\*\*)', content)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                r = p.add_run(part[2:-2])
                r.font.bold = True
            else:
                r = p.add_run(part)
            r.font.name = "Arial"
            r.font.size = Pt(11)
    elif line_s == "---":
        continue
    else:
        add_p(line_s)

out_docx = r"d:\BlindtoBright\docs\NOI_DUNG_DU_AN_GOOGLE_DOCS.docx"
doc.save(out_docx)
print("Xuat file Word Google Docs thanh cong tai:", out_docx)
